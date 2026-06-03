#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Generador de Catalogos en Word
Este script crea documentos Word a partir de datos en Excel
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
import pandas as pd

def generar_catalogos():
      """
          Lee datos de Excel y crea un documento Word por cada fila
              """
      print("=" * 50)
      print("Generador Automatico de Catalogos")
      print("=" * 50)

    # Verifica si existen datos
      archivo_datos = 'datos/productos.xlsx'
      if not os.path.exists(archivo_datos):
                print(f"Error: No se encontro {archivo_datos}")
                print("Por favor, coloca tu archivo Excel en la carpeta 'datos/'")
                return

      # Lee los datos
      print(f"\nLeyendo datos de {archivo_datos}...")
      datos = pd.read_excel(archivo_datos)
      print(f"Se encontraron {len(datos)} registros")

    # Crea carpeta de salida
      os.makedirs('output', exist_ok=True)

    # Genera un catalogo por cada fila
      for idx, row in datos.iterrows():
                doc = Document()
                doc.add_heading(f"Catalogo #{idx+1}", level=0)
                doc.add_paragraph(f"Nombre: {row.get('Nombre', 'N/A')}", style='Normal')
                doc.add_paragraph(f"Precio: ${row.get('Precio', 'N/A')}", style='Normal')
                doc.add_paragraph(f"Descripcion: {row.get('Descripcion', 'N/A')}", style='Normal')

        # Guarda el documento
                nombre_salida = f"output/catalogo_{idx+1:04d}.docx"
                doc.save(nombre_salida)
                print(f"Creado: {nombre_salida}")

      print(f"\n✓ Se crearon {len(datos)} catalogos en la carpeta 'output'")
      print("Listo!\n")

if __name__ == '__main__':
      generar_catalogos()
  

#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Generador de Catálogo Excel Formateado
Lee datos de Excel y crea un nuevo archivo con formato profesional
"""

import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

def generar_catalogo_excel(archivo_entrada='datos/productos.xlsx'):
          """
              Genera un catálogo Excel formateado desde un archivo Excel.
                  Columnas esperadas: Item, Pack, Size, Descripción
                      """

    # Verificar que el archivo existe
          if not os.path.exists(archivo_entrada):
                        print(f"Error: No se encuentra {archivo_entrada}")
                        print("Asegúrate de:")
                        print("1. Crear carpeta 'datos/' en la raíz del proyecto")
                        print("2. Colocar tu archivo Excel como 'productos.xlsx' en esa carpeta")
                        return

          # Leer datos
          print("Leyendo datos...")
          datos = pd.read_excel(archivo_entrada)

    # Verificar columnas requeridas
          columnas_requeridas = ['Item', 'Pack', 'Size', 'Descripción']
          columnas_faltantes = [col for col in columnas_requeridas if col not in datos.columns]

    if columnas_faltantes:
                  print(f"Error: Columnas faltantes: {columnas_faltantes}")
                  print(f"Columnas encontradas: {list(datos.columns)}")
                  return

    # Crear carpeta output si no existe
    os.makedirs('output', exist_ok=True)

    # Crear nuevo workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Catálogo'

    # Definir estilos
    encabezado_font = Font(bold=True, size=12, color="FFFFFF")
    encabezado_fill = PatternFill(start_color="1F4788", end_color="1F4788", fill_type="solid")
    encabezado_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    datos_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # Bordes
    thin_border = Border(
                  left=Side(style='thin'),
                  right=Side(style='thin'),
                  top=Side(style='thin'),
                  bottom=Side(style='thin')
    )

    # Alternar colores para filas
    color_1 = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    color_2 = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")

    # Agregar encabezados
    encabezados = ['Item', 'Pack', 'Size', 'Descripción']
    for col_num, encabezado in enumerate(encabezados, 1):
                  celda = ws.cell(row=1, column=col_num)
                  celda.value = encabezado
                  celda.font = encabezado_font
                  celda.fill = encabezado_fill
                  celda.alignment = encabezado_alignment
                  celda.border = thin_border

    # Establecer ancho de columnas
    ws.column_dimensions['A'].width = 15  # Item
    ws.column_dimensions['B'].width = 12  # Pack
    ws.column_dimensions['C'].width = 12  # Size
    ws.column_dimensions['D'].width = 50  # Descripción

    # Altura de encabezado
    ws.row_dimensions[1].height = 25

    # Agregar datos
    print(f"Procesando {len(datos)} productos...")
    for row_num, (idx, row) in enumerate(datos.iterrows(), 2):
                  # Alternancia de colores
                  fila_color = color_2 if row_num % 2 == 0 else color_1

        valores = [
                          str(row.get('Item', 'N/A')),
                          str(row.get('Pack', 'N/A')),
                          str(row.get('Size', 'N/A')),
                          str(row.get('Descripción', 'N/A'))
        ]

        for col_num, valor in enumerate(valores, 1):
                          celda = ws.cell(row=row_num, column=col_num)
                          celda.value = valor
                          celda.border = thin_border
                          celda.alignment = datos_alignment
                          celda.fill = fila_color

        # Ajustar altura de fila si es necesario
        ws.row_dimensions[row_num].height = 30

    # Congelar primera fila (encabezados)
    ws.freeze_panes = 'A2'

    # Guardar archivo
    fecha = datetime.now().strftime("%d_%m_%Y")
    nombre_excel = f"output/Catalogo_Productos_{fecha}.xlsx"

    print("Generando Excel...")
    wb.save(nombre_excel)

    print(f"\n✅ ¡Catálogo Excel generado exitosamente!")
    print(f"📄 Archivo: {nombre_excel}")
    print(f"📊 Productos incluidos: {len(datos)}")
    print(f"\nPuedes encontrar el archivo en la carpeta 'output/'")

if __name__ == '__main__':
          generar_catalogo_excel()
